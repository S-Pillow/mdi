"""
MDI-54B — XLSX export with result-based row coloring tests.

No live VT/RDAP/WHOIS/DNS/Cloudflare calls. No quota consumed.
"""

from __future__ import annotations

import asyncio
import csv
import io
import sys
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from fastapi.testclient import TestClient

from main import app
from models import (
    DomainHold,
    DomainHoldType,
    DomainStatusSource,
    ExportScope,
    Freshness,
    InputType,
    MdiProcessingStatus,
    MdiResultItem,
    TriageBucket,
    TriageConfidence,
    VtModeUsed,
)
from services import export_service, job_store
from services.export_service import (
    CSV_HEADERS,
    XLSX_RESULT_ROW_FILLS,
    _analyst_row_values,
    build_xlsx_export,
)
from services.job_store import InMemoryJob, clear_all_jobs

client = TestClient(app)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _run(coro):
    return asyncio.run(coro)


def _parse_csv(text: str) -> list[list[str]]:
    return list(csv.reader(io.StringIO(text)))


def _load_xlsx(content: bytes):
    return openpyxl.load_workbook(BytesIO(content), data_only=False)


def _make_job(items: list[MdiResultItem]) -> InMemoryJob:
    return InMemoryJob(
        job_id="mdi54b-test",
        status="done",
        submitted_at=1_700_000_000.0,
        completed_at=1_700_000_100.0,
        total=len(items),
        processed=len(items),
        item_order=[i.input_normalized for i in items],
        results={i.input_normalized: i for i in items},
    )


def _register(job: InMemoryJob) -> None:
    clear_all_jobs()
    _run(job_store.create_job(job))


class TestAnalystRowValues:
    def test_xlsx_ratio_plain_csv_tab_prefixed(self):
        row = MdiResultItem(
            input_original="a.example",
            input_normalized="a.example",
            input_type=InputType.DOMAIN,
            position=1,
            vt_mode_used=VtModeUsed.URL_REPORT,
            vt_hits=1,
            vt_total_engines=91,
            vt_ratio="1/91",
            vt_scan_date=datetime(2026, 6, 1, tzinfo=timezone.utc),
        )
        values = _analyst_row_values(row)
        assert values[7] == "1/91"
        _register(_make_job([row]))
        csv_resp = client.post(
            "/api/mdi/jobs/mdi54b-test/export",
            json={"format": "csv", "scope": "all"},
        )
        csv_rows = _parse_csv(csv_resp.text)
        assert csv_rows[1][7] == "\t1/91"


class TestXlsxContract:
    @pytest.fixture(autouse=True)
    def _clean_jobs(self):
        clear_all_jobs()
        yield
        clear_all_jobs()

    def test_xlsx_http_headers(self):
        row = MdiResultItem(
            input_original="flagged.example",
            input_normalized="flagged.example",
            input_type=InputType.DOMAIN,
            position=1,
            vt_mode_used=VtModeUsed.URL_REPORT,
            vt_hits=2,
            vt_total_engines=91,
            vt_ratio="2/91",
            vt_scan_date=datetime(2026, 6, 1, tzinfo=timezone.utc),
            triage_bucket=TriageBucket.MALICIOUS,
        )
        _register(_make_job([row]))
        resp = client.post(
            "/api/mdi/jobs/mdi54b-test/export",
            json={"format": "xlsx", "scope": "all"},
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(XLSX_MIME)
        cd = resp.headers.get("content-disposition", "")
        assert ".xlsx" in cd

    def test_xlsx_header_and_values_match_csv(self):
        scan_dt = datetime(2026, 6, 1, 15, 30, tzinfo=timezone.utc)
        created = datetime(2020, 3, 15, tzinfo=timezone.utc)
        row = MdiResultItem(
            input_original="flagged.example",
            input_normalized="flagged.example",
            input_type=InputType.DOMAIN,
            position=1,
            vt_mode_used=VtModeUsed.URL_REPORT,
            vt_hits=2,
            vt_suspicious=0,
            vt_total_engines=91,
            vt_ratio="2/91",
            vt_scan_date=scan_dt,
            created_at=created,
            registrar="Example Registrar",
            triage_bucket=TriageBucket.MALICIOUS,
            triage_confidence=TriageConfidence.HIGH,
            cloudflare_rank_available=True,
            cloudflare_rank_bucket="UNRANKED",
            cloudflare_rank_raw_bucket=">200000",
        )
        _register(_make_job([row]))
        csv_resp = client.post(
            "/api/mdi/jobs/mdi54b-test/export",
            json={"format": "csv", "scope": "all"},
        )
        xlsx_resp = client.post(
            "/api/mdi/jobs/mdi54b-test/export",
            json={"format": "xlsx", "scope": "all"},
        )
        csv_rows = _parse_csv(csv_resp.text)
        wb = _load_xlsx(xlsx_resp.content)
        ws = wb.active
        assert [ws.cell(1, c).value for c in range(1, 11)] == CSV_HEADERS
        xlsx_values = [ws.cell(2, c).value for c in range(1, 11)]
        assert xlsx_values[0] == csv_rows[1][0]
        assert xlsx_values[1] == csv_rows[1][1]
        assert xlsx_values[2] == csv_rows[1][2]
        assert csv_rows[1][7] == "\t2/91"
        assert xlsx_values[7] == "2/91"
        assert ws.cell(2, 8).number_format == "@"

    def test_detection_ratio_text_not_date(self):
        row = MdiResultItem(
            input_original="ratio.example",
            input_normalized="ratio.example",
            input_type=InputType.DOMAIN,
            position=1,
            vt_mode_used=VtModeUsed.URL_REPORT,
            vt_hits=1,
            vt_total_engines=91,
            vt_ratio="1/91",
            vt_scan_date=datetime(2026, 6, 1, tzinfo=timezone.utc),
        )
        _register(_make_job([row]))
        resp = client.post(
            "/api/mdi/jobs/mdi54b-test/export",
            json={"format": "xlsx", "scope": "all"},
        )
        ws = _load_xlsx(resp.content).active
        cell = ws.cell(2, 8)
        assert cell.value == "1/91"
        assert cell.number_format == "@"
        assert cell.data_type == "s"

    def test_row_coloring_by_result(self):
        rows = [
            MdiResultItem(
                input_original="flagged.example",
                input_normalized="flagged.example",
                input_type=InputType.DOMAIN,
                position=1,
                vt_mode_used=VtModeUsed.URL_REPORT,
                vt_hits=1,
                vt_total_engines=91,
                vt_ratio="1/91",
                vt_scan_date=datetime(2026, 6, 1, tzinfo=timezone.utc),
                triage_bucket=TriageBucket.MALICIOUS,
            ),
            MdiResultItem(
                input_original="clean.example",
                input_normalized="clean.example",
                input_type=InputType.DOMAIN,
                position=2,
                vt_mode_used=VtModeUsed.URL_REPORT,
                vt_hits=0,
                vt_total_engines=91,
                vt_ratio="0/91",
                vt_scan_date=datetime(2026, 6, 1, tzinfo=timezone.utc),
                triage_bucket=TriageBucket.CLEAN,
            ),
        ]
        _register(_make_job(rows))
        resp = client.post(
            "/api/mdi/jobs/mdi54b-test/export",
            json={"format": "xlsx", "scope": "all"},
        )
        ws = _load_xlsx(resp.content).active
        flagged_fill = ws.cell(2, 1).fill.fgColor.rgb
        clean_fill = ws.cell(3, 1).fill.fgColor.rgb
        assert flagged_fill == XLSX_RESULT_ROW_FILLS["flagged"]
        assert clean_fill == XLSX_RESULT_ROW_FILLS["clean"]
        assert ws.cell(2, 3).value == "flagged"
        assert ws.cell(3, 3).value == "clean"

    def test_export_flagged_scope(self):
        clean = MdiResultItem(
            input_original="clean.example",
            input_normalized="clean.example",
            input_type=InputType.DOMAIN,
            position=1,
            vt_mode_used=VtModeUsed.URL_REPORT,
            vt_hits=0,
            vt_total_engines=91,
            vt_ratio="0/91",
            vt_scan_date=datetime(2026, 6, 1, tzinfo=timezone.utc),
            triage_bucket=TriageBucket.CLEAN,
        )
        bad = clean.model_copy(
            update={
                "input_original": "bad.example",
                "input_normalized": "bad.example",
                "position": 2,
                "vt_hits": 1,
                "vt_ratio": "1/91",
                "triage_bucket": TriageBucket.MALICIOUS,
            }
        )
        _register(_make_job([clean, bad]))
        resp = client.post(
            "/api/mdi/jobs/mdi54b-test/export",
            json={"format": "xlsx", "scope": "flagged"},
        )
        ws = _load_xlsx(resp.content).active
        assert ws.max_row == 2
        assert ws.cell(2, 1).value == "bad.example"

    def test_export_selected_scope(self):
        a = MdiResultItem(
            input_original="one.example",
            input_normalized="one.example",
            input_type=InputType.DOMAIN,
            position=1,
            vt_mode_used=VtModeUsed.URL_REPORT,
            vt_hits=0,
            vt_total_engines=91,
            vt_ratio="0/91",
            vt_scan_date=datetime(2026, 6, 1, tzinfo=timezone.utc),
        )
        b = a.model_copy(
            update={"input_original": "two.example", "input_normalized": "two.example", "position": 2}
        )
        _register(_make_job([a, b]))
        resp = client.post(
            "/api/mdi/jobs/mdi54b-test/export",
            json={"format": "xlsx", "scope": "selected", "positions": [2]},
        )
        ws = _load_xlsx(resp.content).active
        assert ws.max_row == 2
        assert ws.cell(2, 1).value == "two.example"

    def test_on_hold_row_coloring(self):
        hold = DomainHold(
            has_hold=True,
            hold_types=[DomainHoldType.CLIENT_HOLD],
            source=DomainStatusSource.RDAP,
        )
        row = MdiResultItem(
            input_original="held.example",
            input_normalized="held.example",
            input_type=InputType.DOMAIN,
            position=1,
            processing_status=MdiProcessingStatus.HOLD_SKIPPED,
            domain_hold=hold,
            vt_mode_used=VtModeUsed.SKIPPED_HOLD,
            triage_bucket=TriageBucket.HELD,
        )
        _register(_make_job([row]))
        resp = client.post(
            "/api/mdi/jobs/mdi54b-test/export",
            json={"format": "xlsx", "scope": "all"},
        )
        ws = _load_xlsx(resp.content).active
        assert ws.cell(2, 3).value == "on_hold"
        assert ws.cell(2, 1).fill.fgColor.rgb == XLSX_RESULT_ROW_FILLS["on_hold"]

    def test_no_external_calls(self):
        row = MdiResultItem(
            input_original="a.example",
            input_normalized="a.example",
            input_type=InputType.DOMAIN,
            position=1,
            vt_mode_used=VtModeUsed.URL_REPORT,
            vt_hits=0,
            vt_total_engines=91,
            vt_ratio="0/91",
            vt_scan_date=datetime(2026, 6, 1, tzinfo=timezone.utc),
        )
        _register(_make_job([row]))
        with patch("services.vt_service.fetch_report", side_effect=AssertionError("VT called")):
            with patch(
                "services.domain_enrichment.enrich_domain",
                side_effect=AssertionError("RDAP/WHOIS"),
            ):
                with patch(
                    "services.live_check.live_check_item",
                    side_effect=AssertionError("live"),
                ):
                    resp = client.post(
                        "/api/mdi/jobs/mdi54b-test/export",
                        json={"format": "xlsx", "scope": "all"},
                    )
        assert resp.status_code == 200

    def test_json_export_unchanged_schema(self):
        row = MdiResultItem(
            input_original="a.example",
            input_normalized="a.example",
            input_type=InputType.DOMAIN,
            position=1,
            vt_mode_used=VtModeUsed.URL_REPORT,
            vt_targeted_brand="PayPal",
            vt_hits=0,
            vt_total_engines=91,
            vt_ratio="0/91",
            vt_scan_date=datetime(2026, 6, 1, tzinfo=timezone.utc),
        )
        _register(_make_job([row]))
        resp = client.post(
            "/api/mdi/jobs/mdi54b-test/export",
            json={"format": "json", "scope": "all"},
        )
        import json

        data = json.loads(resp.text)
        assert data["schema_version"] == "mdi-evidence-v1"
        assert data["results"][0]["vt_targeted_brand"] == "PayPal"


class TestBuildXlsxExportUnit:
    def test_freeze_and_autofilter(self):
        row = MdiResultItem(
            input_original="a.example",
            input_normalized="a.example",
            input_type=InputType.DOMAIN,
            position=1,
            vt_mode_used=VtModeUsed.URL_REPORT,
            vt_hits=0,
            vt_total_engines=91,
            vt_ratio="0/91",
            vt_scan_date=datetime(2026, 6, 1, tzinfo=timezone.utc),
        )
        job = _make_job([row])
        payload = build_xlsx_export(job, [row], ExportScope.ALL)
        ws = _load_xlsx(payload).active
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref == "A1:J2"
